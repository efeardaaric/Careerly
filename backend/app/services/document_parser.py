"""Document text extraction — PDF/DOCX only. No OCR."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

from docx import Document
from pypdf import PdfReader

ALLOWED_EXTENSIONS = {"pdf", "docx"}
MIN_MEANINGFUL_CHARS = 120


class DocumentParseError(Exception):
    def __init__(self, code: str, message: str):
        self.code = code
        self.message = message
        super().__init__(message)


@dataclass
class ExtractedDocument:
    text: str
    extension: str
    page_count: int | None
    warnings: list[str]


def _sanitize(text: str) -> str:
    # Strip control chars except newline/tab; collapse extreme whitespace.
    cleaned = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", text)
    cleaned = cleaned.replace("\r\n", "\n").replace("\r", "\n")
    cleaned = re.sub(r"[ \t]+\n", "\n", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    return cleaned.strip()


def extract_text(filename: str, data: bytes) -> ExtractedDocument:
    name = filename.lower().strip()
    if "." not in name:
        raise DocumentParseError("invalid_extension", "File must be PDF or DOCX.")
    ext = name.rsplit(".", 1)[-1]
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentParseError("invalid_extension", "Only PDF and DOCX are supported.")

    if not data:
        raise DocumentParseError("empty_file", "Uploaded file is empty.")

    warnings: list[str] = []
    page_count: int | None = None

    try:
        if ext == "pdf":
            reader = PdfReader(io.BytesIO(data))
            page_count = len(reader.pages)
            parts: list[str] = []
            for page in reader.pages:
                parts.append(page.extract_text() or "")
            raw = "\n".join(parts)
        else:
            doc = Document(io.BytesIO(data))
            raw = "\n".join(p.text for p in doc.paragraphs)
            page_count = None
    except Exception as exc:
        raise DocumentParseError(
            "malformed_document",
            "Could not read this document. It may be password-protected or corrupt.",
        ) from exc

    text = _sanitize(raw)
    if len(text) < MIN_MEANINGFUL_CHARS:
        warnings.append("SCANNED_DOCUMENT_REQUIRES_OCR")

    return ExtractedDocument(
        text=text,
        extension=ext,
        page_count=page_count,
        warnings=warnings,
    )
