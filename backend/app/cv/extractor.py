"""PDF (PyMuPDF) and DOCX (python-docx) extraction. No OCR."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path


class CVFileError(Exception):
    def __init__(self, code: str, message: str, status_code: int = 422):
        self.code = code
        self.message = message
        self.status_code = status_code
        super().__init__(message)


@dataclass
class ExtractedDocument:
    text: str
    page_count: int | None
    extraction_method: str
    possibly_scanned: bool = False
    contains_tables: bool = False
    metadata: dict = field(default_factory=dict)


def extract(file_path: str, mime_type: str | None = None) -> ExtractedDocument:
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".pdf" or (mime_type or "").endswith("pdf"):
        return _extract_pdf(path)
    if suffix == ".docx" or "wordprocessingml" in (mime_type or ""):
        return _extract_docx(path)
    raise CVFileError(
        "UNSUPPORTED_CV_TYPE",
        "Please upload a PDF or DOCX file.",
        status_code=415,
    )


def _extract_pdf(path: Path) -> ExtractedDocument:
    import fitz

    try:
        document = fitz.open(path)
    except Exception as exc:
        raise CVFileError(
            "CV_EXTRACTION_FAILED",
            "We couldn't read this CV.",
        ) from exc
    try:
        pages: list[str] = []
        for page in document:
            pages.append(page.get_text("text") or "")
        text = "\n".join(pages)
        page_count = document.page_count
    finally:
        document.close()
    meaningful = len(re.sub(r"\s+", "", text))
    scanned = page_count > 0 and meaningful < 40
    if scanned:
        raise CVFileError(
            "SCANNED_DOCUMENT_DETECTED",
            "This CV appears to contain scanned pages. OCR support is not available yet.",
        )
    if meaningful < 20:
        raise CVFileError(
            "CV_TEXT_EXTRACTION_FAILED",
            "We couldn't read text from this CV.",
        )
    return ExtractedDocument(
        text=text,
        page_count=page_count,
        extraction_method="pymupdf",
        possibly_scanned=False,
        metadata={"character_count": len(text), "word_count": len(text.split())},
    )


def _extract_docx(path: Path) -> ExtractedDocument:
    from docx import Document
    from docx.oxml.ns import qn

    try:
        document = Document(str(path))
    except Exception as exc:
        raise CVFileError(
            "CV_EXTRACTION_FAILED",
            "We couldn't read this CV.",
        ) from exc
    lines: list[str] = []
    contains_tables = False
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            text = "".join(node.text or "" for node in child.iter(qn("w:t"))).strip()
            if text:
                lines.append(text)
        elif child.tag == qn("w:tbl"):
            contains_tables = True
            for row in child.iter(qn("w:tr")):
                cells = []
                for cell in row.iter(qn("w:tc")):
                    cell_text = "".join(
                        node.text or "" for node in cell.iter(qn("w:t"))
                    ).strip()
                    if cell_text:
                        cells.append(cell_text)
                if cells:
                    lines.append(" | ".join(cells))
    text = "\n".join(lines)
    if len(re.sub(r"\s+", "", text)) < 20:
        raise CVFileError(
            "CV_TEXT_EXTRACTION_FAILED",
            "We couldn't read text from this CV.",
        )
    return ExtractedDocument(
        text=text,
        page_count=None,
        extraction_method="python-docx",
        contains_tables=contains_tables,
        metadata={"character_count": len(text), "word_count": len(text.split())},
    )
